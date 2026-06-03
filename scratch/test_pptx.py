from pptx import Presentation
from pptx.util import Inches

def create_test_pptx():
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "Here are some notes"
    prs.save("test_input.pptx")
    print("Created test_input.pptx")

def modify_notes_master():
    prs = Presentation("test_input.pptx")
    # python-pptx has notes_master
    notes_master = prs.notes_master
    print("Found notes master")
    for shape in notes_master.shapes:
        print(f"Shape name: {shape.name}")
        # Move it just to test
        shape.left = Inches(1)
        
    prs.save("test_output.pptx")
    print("Saved test_output.pptx")

create_test_pptx()
modify_notes_master()
